from pathlib import Path
import textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "disclosure_figures"
OUT.mkdir(parents=True, exist_ok=True)

W, H = 1800, 1180
NAVY = "#11263d"; BLUE = "#1d5d8f"; CYAN = "#dff3fb"; GREEN = "#e6f4ea"; AMBER = "#fff4d6"; GRAY = "#edf1f5"; RED = "#fce8e6"
font_paths = ["C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/Arial.ttf"]
bold_paths = ["C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/Arial Bold.ttf"]
def font(size, bold=False):
    for p in (bold_paths if bold else font_paths):
        if Path(p).exists(): return ImageFont.truetype(p, size)
    return ImageFont.load_default()
F12=font(22); F11=font(18); F14=font(25); F16=font(30); F18=font(34,True); F22=font(42,True)

def circuit(path: Path, labels=True):
    im=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(im)
    d.rectangle((25,25,W-25,H-25),outline=NAVY,width=4)
    if labels:
        d.text((60,48),"Edge-AI Multi-Sensor Urban Emergency Detection",font=F22,fill=NAVY)
        d.text((62,100),"Raspberry Pi 4 Model B hardware wiring diagram",font=F16,fill=BLUE)
    # rails
    d.rounded_rectangle((70,175,250,1020),18,fill=RED,outline="#cc5050",width=3)
    if labels:
        d.text((95,200),"5 V RAIL",font=F18,fill="#922b2b")
        d.text((92,245),"Pin 2 / 4",font=F14,fill=NAVY)
    d.rounded_rectangle((270,175,450,1020),18,fill=CYAN,outline=BLUE,width=3)
    if labels:
        d.text((292,200),"3.3 V RAIL",font=F18,fill=BLUE)
        d.text((307,245),"Pin 1 / 17",font=F14,fill=NAVY)
    d.rounded_rectangle((470,175,650,1020),18,fill=GRAY,outline="#718096",width=3)
    if labels:
        d.text((500,200),"COMMON GND",font=F16,fill=NAVY)
        d.text((527,245),"Pins 6, 9, 14...",font=F14,fill=NAVY)
    # controller
    d.rounded_rectangle((710,175,1090,1020),24,fill="#e9eefb",outline=BLUE,width=4)
    if labels:
        d.text((780,205),"Raspberry Pi 4",font=F22,fill=NAVY)
        d.text((790,255),"40-pin header",font=F16,fill=BLUE)
    pin_rows=[("Pin 1", "3V3"),("Pin 3", "GPIO2 / SDA1"),("Pin 5", "GPIO3 / SCL1"),("Pin 7", "GPIO4"),("Pin 8", "GPIO14 / TX"),("Pin 10", "GPIO15 / RX"),("Pin 12", "GPIO18 / I2S BCLK"),("Pin 29", "GPIO5 / RED"),("Pin 31", "GPIO6 / YELLOW"),("Pin 32", "GPIO12 / PWM0"),("Pin 33", "GPIO13 / GREEN"),("Pin 35", "GPIO19 / I2S WS"),("Pin 36", "GPIO16 / BUZZER"),("Pin 37", "GPIO26 UNUSED"),("Pin 38", "GPIO20 / I2S DIN")]
    y=330
    for p,n in pin_rows:
        fill="#f8fbff" if "UNUSED" not in n else "#ffe9e9"
        d.rounded_rectangle((745,y,1055,y+35),8,fill=fill,outline="#b4c5d8")
        if labels:
            d.text((760,y+8),p,font=F11,fill=NAVY); d.text((875,y+8),n,font=F11,fill="#9b2c2c" if "UNUSED" in n else NAVY)
        y+=42
    # component cards
    comps=[
      (1160,180,"DHT22", "VCC→3V3; DATA→GPIO4 Pin 7; GND→GND", CYAN),
      (1160,295,"GY-87 / MPU6050", "3V3; GND; SDA→GPIO2 Pin 3; SCL→GPIO3 Pin 5", CYAN),
      (1160,410,"ADS1115 + MQ-2", "VDD→3V3; I2C; ADDR→GND; A0←MQ-2 AO", AMBER),
      (1160,525,"INMP441", "VDD→3V3; GND; SCK→GPIO18; WS→GPIO19; SD→GPIO20; L/R→GND", CYAN),
      (1160,640,"NEO-6M GPS", "TX→GPIO15 Pin 10; RX→GPIO14 Pin 8; VCC appropriate; GND", GREEN),
      (1160,755,"Traffic LEDs", "R→GPIO5 Pin 29; Y→GPIO6 Pin 31; G→GPIO13 Pin 33; each 220–330Ω", GREEN),
      (1160,870,"Buzzer + SG90", "Buzzer I/O→GPIO16 Pin 36 PWM; Servo signal→GPIO12 Pin 32; servo VCC→5V", RED),
    ]
    for x,y,title,body,fill in comps:
        d.rounded_rectangle((x,y,1715,y+100),16,fill=fill,outline=BLUE,width=3)
        if labels:
            d.text((x+20,y+12),title,font=F18,fill=NAVY)
            wrapped="\n".join(textwrap.wrap(body,width=38,break_long_words=True,break_on_hyphens=False))
            d.multiline_text((x+20,y+52),wrapped,font=F12,fill=NAVY,spacing=2)
    # wires / bus hints
    d.line((650,260,710,260),fill="#c53a3a",width=8)
    d.line((650,300,710,300),fill=BLUE,width=8)
    d.line((650,340,710,340),fill="#4a5568",width=8)
    if labels:
        d.text((520,265),"3.3V",font=F12,fill="#922b2b")
        d.text((520,305),"I2C / UART / I2S / GPIO",font=F12,fill=BLUE)
        d.text((520,345),"COMMON GND",font=F12,fill=NAVY)
        d.text((75,1060),"Safety: MQ-2 AO must pass through a divider before ADS1115 A0; never connect it to a Pi GPIO. SG90 uses a suitable 5 V supply with common ground.",font=F14,fill="#7b341e")
        d.text((75,1100),"USB webcam: direct USB port. GPIO26 / physical pin 37: UNUSED.",font=F14,fill=NAVY)
    im.save(path,dpi=(180,180))

def architecture(path: Path):
    im=Image.new("RGB",(W,900),"white"); d=ImageDraw.Draw(im); d.rectangle((25,25,W-25,875),outline=NAVY,width=4)
    d.text((65,55),"SENTINEL-AI operating architecture and emergency response flow",font=F22,fill=NAVY)
    boxes=[(70,180,430,385,"SENSORS","DHT22 • GY-87 • MQ-2\nINMP441 • GPS • camera",CYAN),(460,180,820,385,"EDGE INFERENCE","Audio + vision models\nTemporal fusion and OOD",CYAN),(850,180,1210,385,"GOVERNED DECISION","Risk and assurance\nSafety contract",AMBER),(1240,180,1730,385,"OPERATOR WORKFLOW","Acknowledge and confirm\nEscalate and resolve",GREEN)]
    for x1,y1,x2,y2,title,body,fill in boxes:
        d.rounded_rectangle((x1,y1,x2,y2),20,fill=fill,outline=BLUE,width=4); d.text((x1+25,y1+25),title,font=F18,fill=NAVY); d.multiline_text((x1+25,y1+95),body,font=F14,fill=NAVY,spacing=10)
    for a,b in [((390,275),(475,275)),((785,275),(870,275)),((1175,275),(1260,275))]:
        d.line((*a,*b),fill=BLUE,width=8); d.polygon([(b[0],b[1]),(b[0]-18,b[1]-12),(b[0]-18,b[1]+12)],fill=BLUE)
    d.rounded_rectangle((300,520,1500,760),20,fill=GRAY,outline="#718096",width=4)
    d.text((350,555),"AUDITABLE OFFLINE-FIRST RECORD",font=F18,fill=NAVY)
    d.multiline_text((350,625),"Incident timeline • evidence files • model version • OOD reasons • operator actions\nSHA-256 verification • local SQLite/WAL • outbox sync when connectivity returns",font=F16,fill=NAVY,spacing=14)
    d.line((895,370,895,520),fill=BLUE,width=8); d.polygon([(895,520),(883,495),(907,495)],fill=BLUE)
    im.save(path,dpi=(180,180))

def build_docx():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    styles=doc.styles; styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(9)
    p=doc.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('SENTINEL AI Disclosure Figures')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Supplementary technical drawings for the Edge AI urban emergency detection prototype')
    doc.add_paragraph('The existing project circuit diagram is intentionally preserved in the original disclosure document. This appendix supplies the additional labeled and unlabeled figure sets requested for the complete specification.')
    for title,img,caption in [('Figure 2. Labeled hardware wiring diagram',OUT/'hardware_wiring_labeled.png','Figure 2: Edge-AI Multi-Sensor Urban Emergency Detection – Raspberry Pi 4 Hardware Wiring Diagram'),('Figure 3. System architecture and response flow',OUT/'system_architecture.png','Figure 3: Multimodal sensing, governed decision, operator workflow and audit record')]:
        doc.add_heading(title,level=1); p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(img),width=Inches(7.0)); p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading('Connection table',level=1)
    rows=[('DHT22','VCC / DATA / GND','1 / 7 / 9','3.3V / GPIO4 / GND','Temperature and humidity'),('GY-87 / MPU6050','3.3V / GND / SDA / SCL','1 / 6 / 3 / 5','3.3V / GND / GPIO2 / GPIO3','Impact and motion over I2C'),('ADS1115','VDD / GND / SDA / SCL / ADDR / A0','1 / 6 / 3 / 5 / GND / ADC','3.3V / GND / GPIO2 / GPIO3 / 0x48 / A0','ADC for MQ-2; use divider'),('MQ-2','VCC / GND / AO / DO','5V rail / GND / ADS1115 A0 / NC','5V / GND / analog','Smoke and gas; never direct to GPIO'),('INMP441','VDD / GND / SCK / WS / SD / L/R','1 / GND / 12 / 35 / 38 / GND','3.3V / GPIO18 / GPIO19 / GPIO20','I2S audio'),('NEO-6M','VCC / GND / TX / RX','appropriate / GND / 10 / 8','GPIO15 RX / GPIO14 TX','UART position'),('LEDs','R / Y / G anodes; cathodes','29 / 31 / 33 / GND','GPIO5 / GPIO6 / GPIO13','Each anode through 220–330Ω'),('Buzzer','VCC / GND / I/O','3.3V / GND / 36','GPIO16 PWM','Active-high or PWM module'),('SG90','VCC / GND / signal','5V / common GND / 32','GPIO12 PWM0','Separate suitable 5V supply recommended'),('USB webcam','USB','USB port','No GPIO','Video input'),('Reserved','GPIO26','37','UNUSED','Must remain unconnected')]
    table=doc.add_table(rows=1,cols=5); table.style='Table Grid'; hdr=table.rows[0].cells
    for i,h in enumerate(['Component','Component pin','Physical pin','GPIO or signal','Purpose']): hdr[i].text=h
    for row in rows:
        cells=table.add_row().cells
        for i,v in enumerate(row): cells[i].text=v
    doc.add_heading('Electrical notes',level=1)
    doc.add_paragraph('The Raspberry Pi GPIO header is 3.3 V logic. The MQ-2 heater is powered from 5 V and its analog output must be reduced with a suitable voltage divider before entering the 3.3 V ADS1115 input. The SG90 servo must use a suitable 5 V supply and share ground with the Raspberry Pi; its motor current should not be drawn through a GPIO pin. The drawing uses physical 40-pin header numbering and BCM GPIO names together.')
    doc.save(OUT/'SENTINEL_AI_Disclosure_Figures.docx')

def main():
    circuit(OUT/'hardware_wiring_labeled.png',True); circuit(OUT/'hardware_wiring_clean.png',False); architecture(OUT/'system_architecture.png'); build_docx()
if __name__=='__main__': main()
