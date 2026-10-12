from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

root = Path(r'C:\Users\gokul\Desktop\PROJECTS\Edge-AI')
out = root / 'docs/disclosure_figures'
source = Path(r'C:\Users\gokul\Downloads\SENTINEL_AI_Invention_Disclosure_Final.docx')
photo = Path(r'C:\Users\gokul\AppData\Local\Temp\codex-clipboard-0cbb0dd4-ad7a-4375-8aab-9cb0a122af80.jpg')
im = Image.new('RGB', (1800, 1160), 'white')
draw = ImageDraw.Draw(im)
def font(size): return ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', size)
def text(x,y,s,size=26): draw.text((x,y),s,font=font(size),fill='black')
def box(x,y,w,h,lines):
    draw.rectangle((x,y,x+w,y+h),outline='black',width=3)
    for i,line in enumerate(lines):
        f=font(26); length=draw.textlength(line,font=f)
        assert length < w-20, line
        draw.text((x+(w-length)/2,y+20+i*36),line,font=f,fill='black')
def arrow(x,y,xx,yy):
    draw.line((x,y,xx,yy),fill='black',width=3)
    if yy>y: draw.polygon([(xx,yy),(xx-9,yy-16),(xx+9,yy-16)],fill='black')
    else: draw.polygon([(xx,yy),(xx-16,yy-9),(xx-16,yy+9)],fill='black')
text(70,35,'SENTINEL-AI working principle',38)
box(60,120,460,165,['Sound and physical observations','INMP441 microphone; sensors','GPS location; camera input'])
box(650,120,490,165,['Raspberry Pi 4 edge processing','Training-matched preprocessing','Audio / vision inference'])
box(1270,120,460,165,['Reliability and evidence checks','Sensor health; model assurance','Temporal agreement; conflicts'])
arrow(520,202,650,202); arrow(1140,202,1270,202)
arrow(1500,285,1500,380)
box(1270,380,460,165,['Response plan and authorization','Apply safety contract and roles','Uncertain evidence: human review'])
box(60,680,460,160,['Traffic LED indication','Red / yellow / green','Display permitted traffic state'])
box(650,680,490,160,['SG90 lane-barrier actuation','Authorized command only','Confirm position / outcome'])
box(1270,680,460,160,['Operator interface','Review, acknowledge and resolve','External dispatch: human approval'])
draw.line((1500,545,1500,605,290,605),fill='black',width=3)
arrow(290,605,290,680); arrow(895,605,895,680); arrow(1500,605,1500,680)
text(660,565,'Permitted response commands',24)
box(260,960,1280,140,['Local incident record: predictions, timestamps, evidence hashes and operator actions',
                        'Offline queue retains records for synchronization when connectivity returns'])
for x in (290,895,1500): arrow(x,840,x,960)
text(70,1120,'Functional line diagram; physical performance and actuation require hardware validation.',23)
fig = out / 'sentinel_working_principle_fig3.png'
im.save(fig)
d=Document(source)
cell=d.tables[0].rows[24].cells[-1]
intro=cell.paragraphs[0]
intro.text='The following figures present the labelled SENTINEL-AI prototype (Figure 1), its system architecture (Figure 2), and the functional line diagram explaining its working principle (Figure 3).'
arch=next(p for p in cell.paragraphs if p._p.xpath('.//w:drawing'))
def picture_before(anchor,path,width):
    p=anchor.insert_paragraph_before(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Inches(width))
    return p
def caption_before(anchor,value):
    p=anchor.insert_paragraph_before(value); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(10)
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(11)
    return p
picture_before(arch,photo,6.2)
caption_before(arch,'Fig. 1: Edge-AI Multi-Sensor Urban Emergency Detection, Localization & Autonomous Response — Prototype')
anchor=next(p for p in cell.paragraphs if p.text=='Key connection summary')
picture_before(anchor,fig,6.4)
caption_before(anchor,'Fig. 3: Line diagram illustrating the SENTINEL-AI working principle')
output=out/'SENTINEL_AI_Disclosure_With_Prototype.docx'
d.save(output)
print(output)
print('Figures in Section 12:',[p.text for p in cell.paragraphs if p.text.startswith('Fig.')])
