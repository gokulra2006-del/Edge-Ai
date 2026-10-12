from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\gokul\Downloads\Invention_Disclosure_Form_SENTINEL_AI_Project_Enhanced.docx")
OUT = ROOT / "docs" / "disclosure_figures" / "SENTINEL_AI_Invention_Disclosure_Complete.docx"
FIG = ROOT / "docs" / "disclosure_figures"

def set_cell(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Arial"; r.font.size = Pt(9)

def add_figure(cell, path, caption):
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.4))
    q = cell.add_paragraph(caption); q.alignment = WD_ALIGN_PARAGRAPH.CENTER

def main():
    doc = Document(str(SOURCE)); t = doc.tables[0]
    content = {
      1: "Edge-AI Multi-Sensor Urban Emergency Detection, Localization and Governed Response System",
      3: "The invention is an offline-first edge-computing safety system for urban intersections. It addresses delayed and unreliable emergency detection caused by single-camera systems, isolated alarms, cloud dependence and unverified automatic actuation. A Raspberry Pi 4 combines acoustic, visual, thermal/humidity, gas/smoke, inertial and GPS evidence; produces a localized risk assessment; presents a governed response plan; and records an auditable incident package. Emergency-service contact remains operator-approved and the system never silently claims that a synthetic model is field validated.",
      5: "Edge AI; Raspberry Pi 4; multimodal sensor fusion; urban emergency detection; crash detection; fire and smoke detection; acoustic event classification; INMP441; GY-87 MPU6050; MQ-2 gas sensor; ADS1115 ADC; DHT22; NEO-6M GPS; localization; traffic-corridor response; operator-in-the-loop; assurance; out-of-distribution detection; offline-first incident evidence; tamper-evident audit log; black-box recorder.",
      8: "1. The sensors acquire synchronized observations at the edge. 2. Audio, camera and telemetry pipelines produce event candidates with timestamps and model identifiers. 3. A temporal fusion and deep-rule layer checks corroboration, signal quality, sensor availability and out-of-distribution conditions. 4. The governed decision layer assigns event type, zone, confidence and severity, and generates a response plan. 5. The operator reviews, acknowledges, confirms or rejects the plan; emergency-service notification is never automatic. 6. Approved traffic actions are sent to the actuator layer, including LEDs, buzzer, servo and configured corridor signals. 7. Incidents, predictions, evidence, operator actions, assurance state and outcomes are persisted locally through SQLite/WAL and exported as a hash-verified evidence package. 8. When connectivity returns, an outbox can synchronize records without blocking inference.",
      9: "The inventive contribution is the governed combination of: (i) heterogeneous low-cost sensors on a Raspberry Pi edge node; (ii) temporal, cross-modal corroboration with sensor-availability awareness; (iii) assurance states, model identity, usage restrictions and OOD reasons attached to every prediction; (iv) an operator-in-the-loop response contract that separates recommendation from authorization; (v) offline-first, tamper-evident incident evidence containing timeline, model versions, actions and file hashes; and (vi) safe actuator control with explicit power-domain and GPIO constraints. The novelty lies in this auditable arrangement and decision procedure rather than in claiming any individual sensor or neural network in isolation.",
      11: "Alternative embodiments include additional camera nodes, LoRa/Wi-Fi/Ethernet backhaul, other I2S microphones, calibrated gas sensors, thermal cameras, CAN/vehicle telemetry, a larger ADC, a cloud mirror, or a different edge computer. The safety contract, operator approval, evidence manifest and model-usage restriction remain applicable. The Raspberry Pi embodiment is preferred for the disclosed prototype.",
      13: "I designed SENTINEL-AI so that an emergency is treated as a sequence of evidence and decisions rather than as one alarm. The device senses locally, checks whether different sensors agree, explains why it raised a risk, lets an authorized human approve the response, and preserves what happened for later review. This makes the prototype useful even when the network is unavailable and makes its limitations visible.",
      15: "Prototype validation includes unit and integration tests for the database migrations, incident IDs and merge windows, duplicate suppression, writer-queue resilience, evidence export and tamper detection, model checksum mismatch, role permissions and regression behavior. Simulation live-streams produce incident records and evidence manifests. Results are laboratory/prototype evidence; field performance, emergency dispatch integration and safety certification remain future validation steps.",
      17: "The project uses acoustic, visual and sensor telemetry data. Acoustic classes include ambient, crash, horn and siren; visual classes include fire, smoke and vehicle categories; telemetry includes temperature, humidity, inertial motion, gas level and GPS. Synthetic samples may be used for research augmentation, but the registry marks the profile RESEARCH_ONLY and stores unavailable evaluation metrics as UNVERIFIED. Every stored prediction carries model ID/version, assurance state and usage restriction.",
      20: "Advantages include low-latency local operation; resilience during network loss; multi-modal corroboration that can reduce single-sensor false alarms; explicit uncertainty and OOD handling; operator-controlled escalation; physical pin and power safety; synchronized incident timelines; hash-verified evidence; role-based permissions; and a hardware/software architecture that can be reproduced on low-cost edge hardware.",
      23: "Illustrative prototype ranges are: 3.3 V logic for Raspberry Pi GPIO/I2C/I2S; 5 V supply for MQ-2 heater and SG90 servo; DHT22 data on GPIO4; I2C at GPIO2/GPIO3; I2S on GPIO18/GPIO19/GPIO20; UART on GPIO14/GPIO15; LEDs on GPIO5/GPIO6/GPIO13 with 220–330 ohm resistors; buzzer PWM on GPIO16; servo PWM on GPIO12; ADS1115 address 0x48; MQ-2 AO through a divider into ADS1115 A0; GPIO26/physical pin 37 unused. Actual thresholds are configuration values validated during commissioning.",
      27: "Potential users are traffic control rooms, municipal emergency operations centres, campus/industrial safety teams, road operators, researchers and trained maintenance engineers. Direct operators are trained adults with role-specific permissions: Commander (full incident/system authority), Operator (traffic and acknowledgement), Engineer (sensors, calibration and models), and Viewer (read-only monitoring). Benefits include faster local assessment, explainable response, reduced dependence on cloud connectivity and recoverable evidence.",
      29: "Additional notes: the supplied first/project circuit diagram is intentionally preserved. This completed disclosure adds the faculty-style supplementary labeled schematic, clean drawing set, system architecture flow, connection table and technical notes. The MQ-2 is not connected directly to a Raspberry Pi GPIO; the SG90 has a suitable 5 V source and common ground; USB camera uses a USB port; and GY-87 auxiliary pins VCC_IN, FSYNC, INTA and DRDY remain unconnected unless the selected module requires them. Patentability, inventorship and legal claim scope should be reviewed by the institution's IP office.",
    }
    # Locate content rows by their first-column section number; merged templates use blank labels on content rows.
    for i, row in enumerate(t.rows):
        label = row.cells[0].text.strip().rstrip('.')
        try: key = int(label)
        except ValueError: key = None
        if key in content and i + 1 < len(t.rows):
            set_cell(t.rows[i + 1].cells[1], content[key])
    # Replace the drawing section content while preserving the original diagram in the source document.
    draw_row = t.rows[24]
    set_cell(draw_row.cells[1], "The following supplementary drawings are provided in two sets as requested: a fully labeled set for technical specification and a clean set without descriptive labels for formal drawing use. The original project diagram remains preserved in the source disclosure.")
    add_figure(draw_row.cells[1], FIG/'hardware_wiring_labeled.png', 'Fig. 2: Fully labeled Raspberry Pi 4 hardware wiring diagram')
    add_figure(draw_row.cells[1], FIG/'hardware_wiring_clean.png', 'Fig. 3: Clean wiring drawing without descriptive labels')
    add_figure(draw_row.cells[1], FIG/'system_architecture.png', 'Fig. 4: System architecture and emergency response flow')
    # Add a compact connection table to the drawing section.
    p = draw_row.cells[1].add_paragraph(); p.add_run('Key connection summary').bold = True
    table = draw_row.cells[1].add_table(rows=1, cols=4); table.style = 'Table Grid'
    for c, h in zip(table.rows[0].cells, ['Component','Pin / signal','Pi physical pin','GPIO / purpose']): c.text = h
    data = [('DHT22','DATA','7','GPIO4 / temperature-humidity'),('GY-87','SDA, SCL','3, 5','GPIO2, GPIO3 / I2C'),('ADS1115','SDA, SCL, ADDR','3, 5, GND','0x48 / ADC'),('MQ-2','AO','ADS1115 A0','5 V sensor; divider required'),('INMP441','SCK, WS, SD','12, 35, 38','GPIO18, 19, 20 / I2S'),('GPS','TX, RX','10, 8','GPIO15 RX, GPIO14 TX'),('LEDs','R, Y, G','29, 31, 33','GPIO5, 6, 13'),('Buzzer / servo','I/O / signal','36 / 32','GPIO16 PWM / GPIO12 PWM'),('Reserved','GPIO26','37','UNUSED')]
    for vals in data:
        cells = table.add_row().cells
        for c, v in zip(cells, vals): c.text = v
    doc.core_properties.title = 'SENTINEL-AI Invention Disclosure – Complete Technical Version'
    doc.core_properties.subject = 'Edge-AI multi-sensor emergency detection, localization and governed response'
    doc.save(str(OUT))
    print(OUT)

if __name__ == '__main__': main()
